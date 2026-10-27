"""키워드 원본 엑셀 + 추가 키워드 → data/site-pages.json (사이트 라우트의 단일 출처) + keywords/누수탐지_일산_510.xlsx

  python -I scripts/build_pages.py

오픈 단계(wave)
  1 = 시·구 핵심 (일산·일산동구·일산서구·김포·파주·은평구·마포 + 표기 변형) + 시·구 단위 추가 키워드
  2 = 생활권·신도시 (탄현·주엽·한강신도시·운정·연신내 …)
  3 = 동·읍면 (마두동·김포 대곶면·파주 문산읍·은평구 갈현동 …)
"""
import json, sys
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))
from places_src import ALIAS, places  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'source' / 'keywords_src.xlsx'
OUT_XLSX = ROOT / 'keywords' / '누수탐지_일산_510.xlsx'
SITE_URL = 'https://leak.example'  # TODO: 도메인 확정 후 lib/site.ts 와 함께 바꿀 것
TOTAL = 510

# 원본 481개에 더하는 연관 서비스 키워드 (홈페이지 '주요업무' 범위: 누수탐지·누수공사·배관·하수구 막힘·고압세척)
# (키워드, 지역 표기, 서비스, 수식어)
ADDED = [
    ('일산 누수공사', '일산', '누수공사', None),
    ('일산 누수탐지업체', '일산', '누수탐지', '업체'),
    ('일산 배관누수', '일산', '배관누수', None),
    ('일산 아파트 누수', '일산', '아파트 누수', None),
    ('일산 화장실 누수', '일산', '화장실 누수', None),
    ('일산 천장 누수', '일산', '천장 누수', None),
    ('일산 난방배관 누수', '일산', '난방배관 누수', None),
    ('일산 수도 누수', '일산', '수도 누수', None),
    ('일산 하수구막힘', '일산', '하수구막힘', None),
    ('일산 고압세척', '일산', '고압세척', None),
    ('일산동구 누수공사', '일산동구', '누수공사', None),
    ('일산서구 누수공사', '일산서구', '누수공사', None),
    ('김포 누수탐지업체', '김포', '누수탐지', '업체'),
    ('김포 누수공사', '김포', '누수공사', None),
    ('김포 아파트 누수', '김포', '아파트 누수', None),
    ('김포 화장실 누수', '김포', '화장실 누수', None),
    ('김포 하수구막힘', '김포', '하수구막힘', None),
    ('파주 누수탐지업체', '파주', '누수탐지', '업체'),
    ('파주 누수공사', '파주', '누수공사', None),
    ('파주 아파트 누수', '파주', '아파트 누수', None),
    ('파주 천장 누수', '파주', '천장 누수', None),
    ('파주 하수구막힘', '파주', '하수구막힘', None),
    ('은평구 누수탐지업체', '은평구', '누수탐지', '업체'),
    ('은평구 누수공사', '은평구', '누수공사', None),
    ('은평구 빌라 누수', '은평구', '빌라 누수', None),
    ('은평구 하수구막힘', '은평구', '하수구막힘', None),
    ('마포 누수공사', '마포', '누수공사', None),
    ('운정 누수공사', '운정', '누수공사', None),
    ('한강신도시 누수공사', '한강신도시', '누수공사', None),
]

rows = []
for r in list(load_workbook(SRC, read_only=True).worksheets[0].iter_rows(values_only=True))[1:]:
    no, kw, joined, prio, _cls, city, label, service, mod = r[:9]
    if kw:
        rows.append(dict(kw=kw, joined=joined, prio=prio, city=city, label=label, service=service, mod=mod, added=False))
assert len(rows) == 481, len(rows)
for kw, label, service, mod in ADDED:
    key = ALIAS.get(label, label)
    rows.append(dict(kw=kw, joined=kw.replace(' ', ''), prio=2 if places[key]['kind'] == '시구' else 3,
                     city=places[key]['city'], label=label, service=service, mod=mod, added=True))
assert len(rows) == TOTAL, len(rows)

pages, rep = [], {}
for i, r in enumerate(rows, 1):
    key = ALIAS.get(r['label'], r['label'])
    assert key in places, f"지역 데이터 없음: {r['label']} ({r['kw']})"
    kind = places[key]['kind']
    if r['added']:
        wave = 1 if kind == '시구' else 2
    else:
        wave = 1 if (r['prio'] in (1, 2) and kind == '시구') else (2 if r['prio'] in (1, 2) else 3)
    pg = dict(n=i, kw=r['kw'], slug=r['kw'].replace(' ', '-'), label=r['label'], place=key, city=r['city'],
              service=r['service'], mod=r['mod'] or '', prio=r['prio'], wave=wave, added=r['added'])
    k = (key, r['service'], r['mod'] or '')
    if k in rep:
        pg['variantOf'] = rep[k]
    else:
        rep[k] = pg['slug']
    pages.append(pg)

slugs = Counter(p['slug'] for p in pages)
assert all(v == 1 for v in slugs.values()), [k for k, v in slugs.items() if v > 1]
json.dump({'pages': pages}, open(ROOT / 'data' / 'site-pages.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)

# ---- 키워드 엑셀 (첫 시트 = 키워드 시트, 요약 시트 없음) ----
wb = Workbook()
ws = wb.active
ws.title = '키워드'
head = ['No', '키워드', '붙여쓰기', '우선순위', '지역', '서비스', '수식어', 'URL', '오픈 단계', '추가 여부']
ws.append(head)
for p, r in zip(pages, rows):
    ws.append([p['n'], p['kw'], r['joined'], r['prio'], p['city'], p['service'], p['mod'] or '',
               f"{SITE_URL}/{p['slug']}/", p['wave'], '추가' if p['added'] else ''])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='0B2236')
    c.alignment = Alignment(horizontal='center')
for col, w in zip('ABCDEFGHIJ', [6, 30, 30, 9, 9, 14, 10, 52, 9, 9]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:J{len(pages) + 1}'
OUT_XLSX.parent.mkdir(exist_ok=True)
wb.save(OUT_XLSX)

sys.stdout.reconfigure(encoding='utf-8')
c = Counter(p['wave'] for p in pages)
acc = 0
for w in sorted(c):
    acc += c[w]
    print(f'wave {w}: {c[w]} (누적 {acc})')
print('표기 변형(variantOf):', sum(1 for p in pages if 'variantOf' in p), '/ 대표 페이지:', sum(1 for p in pages if 'variantOf' not in p))
print('엑셀:', OUT_XLSX.name, len(pages), '행')
